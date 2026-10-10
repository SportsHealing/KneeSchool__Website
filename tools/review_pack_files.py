"""The spreadsheet and the Word version of a review pack.

A consultant filling 85 cells wants a column they can tab down, not a markdown
table. Both formats carry the same rows as the markdown pack, the same row
references, and both read back through tools/review_pack.py --ingest. That is the
requirement that chose the libraries: openpyxl and python-docx can both read a
file back, and the pack is useless if a reviewer's answers cannot return.

Kept in its own module so tools/review_pack.py stays importable with nothing but
the standard library. The markdown pack, --status and --check never touch this
file, so a checkout without openpyxl installed still generates, checks and
ingests the markdown packs.
"""

import os

# Arial throughout, because a review pack is a document someone prints.
FONT = "Arial"

POSITION_STATES = ["CONFIRMED", "AMENDED", "WITHDRAWN", "NOT_CLINICAL"]
FIGURE_STATES = ["VERIFIED", "CORRECTED", "REMOVED"]

READ_ME = [
    ("What this is",
     "One chapter of KneeSchool, with every measurement and every clinical "
     "recommendation on its pages listed for review. Generated from the pages "
     "themselves by tools/review_pack.py, so this file cannot disagree with the "
     "site."),
    ("What has not happened",
     "No source was retrieved for any page in this chapter. The build "
     "environment has no access to PubMed, Cochrane, a journal or a textbook, so "
     "the pipeline's fact checking stage could not run. Every claim here has been "
     "checked by nobody. This pack stands in for the verification stage as well "
     "as for the review stage."),
    ("What to fill in",
     "The Decision column on the Recommendations and Measurements sheets, and "
     "your name. Every other cell is rewritten when the pack is regenerated. A "
     "row left blank stays outstanding and comes back in the next pack, so "
     "answering part of a sheet is a normal way to use it."),
    ("Recommendations: allowed decisions",
     "CONFIRMED, the sentence stands as written. AMENDED, it needs changing and "
     "the new wording goes in the Note column. WITHDRAWN, take it off the page. "
     "NOT_CLINICAL, it is teaching rather than clinical direction. The detector "
     "over collects on purpose, so NOT_CLINICAL is expected and is not a "
     "complaint about the page."),
    ("Measurements: allowed decisions",
     "VERIFIED, the figure is right and the source goes in the Source column. "
     "CORRECTED, the right value goes in the Note column. REMOVED, take the "
     "figure off the page."),
    ("How to return it",
     "Send the file back. It is read into the registers with: python3 "
     "tools/review_pack.py --ingest <this file>. One malformed row stops the "
     "whole file, so nothing is written half way."),
    ("What this cannot do",
     "It does not change a page. A confirmed recommendation needs nothing; an "
     "amended one needs the prose edited, and your wording in the Note column is "
     "what the editor works from."),
]


def _facts_rows(meta):
    return [
        ("Section", meta["section"]),
        ("Chapter", "%s %s" % (meta["chapter"], meta["name"])),
        ("Pages", meta["pages"]),
        ("Page type", meta["types"]),
        ("Tiers", meta["tiers"]),
        ("Words", meta["words"]),
        ("Style gate", meta["gate"]),
        ("Evidence verification", "did not run on any page"),
        ("Measurements to check", len(meta["unchecked"])),
        ("Recommendations to sign", len(meta["unsigned"])),
    ]


def write_xlsx(path, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    head = Font(name=FONT, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="16382D")      # --green-ink
    body = Font(name=FONT)
    label = Font(name=FONT, bold=True)
    example = Font(name=FONT, italic=True, color="8A6A2F")  # --gold-ink
    example_fill = PatternFill("solid", fgColor="FAF7EF")   # --ivory
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    wb = Workbook()

    def header(ws, titles, widths):
        ws.append(titles)
        for i, (t, w) in enumerate(zip(titles, widths), 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill, c.alignment = head, head_fill, wrap
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    # Read me
    ws = wb.active
    ws.title = "Read me"
    ws.append(["Consultant review pack: chapter %s %s" % (meta["chapter"], meta["name"])])
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws.append([])
    for k, v in _facts_rows(meta):
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = label
        ws.cell(row=ws.max_row, column=2).font = body
    ws.append([])
    for k, v in READ_ME:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = label
        c = ws.cell(row=ws.max_row, column=2)
        c.font, c.alignment = body, wrap
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 96

    # Pages
    ws = wb.create_sheet("Pages")
    header(ws, ["Page", "Title", "Tiers", "Words", "Published as"],
           [10, 40, 8, 10, 52])
    for f in meta["facts"]:
        ws.append([f["page_id"], f["title"], len(f["tiers"]), f["words"],
                   meta["published"].get(f["page_id"], "")])
        for i in range(1, 6):
            ws.cell(row=ws.max_row, column=i).font = body
            ws.cell(row=ws.max_row, column=i).alignment = top

    # Claims
    ws = wb.create_sheet("Claims")
    header(ws, ["Page", "Title", "Claim"], [10, 34, 110])
    for f in meta["facts"]:
        for claim in f["claims"]:
            ws.append([f["page_id"], f["title"], claim])
            for i in range(1, 4):
                ws.cell(row=ws.max_row, column=i).font = body
                ws.cell(row=ws.max_row, column=i).alignment = wrap

    # Recommendations. Column order matches the markdown pack's table, with Note
    # added, so a reviewer who has seen one recognises the other.
    ws = wb.create_sheet("Recommendations")
    header(ws, ["Ref", "Page", "Tier", "Sentence", "Decision", "Reviewer", "Note"],
           [12, 10, 16, 86, 16, 22, 50])
    ws.append(["EXAMPLE", "0.0.0", "mrcs",
               "This row shows the expected format and is ignored when the pack is read back.",
               "CONFIRMED", "Miss A Surgeon", "Leave the Note column empty unless amending."])
    for i in range(1, 8):
        c = ws.cell(row=2, column=i)
        c.font, c.fill, c.alignment = example, example_fill, wrap
    first = 3
    for f, i, p in meta["unsigned"]:
        ws.append(["%s-P%d" % (f["page_id"], i), f["page_id"], p.get("tier", ""),
                   p.get("as_written", ""), "", "", ""])
        for col in range(1, 8):
            ws.cell(row=ws.max_row, column=col).font = body
            ws.cell(row=ws.max_row, column=col).alignment = wrap
    last_pos = ws.max_row
    if meta["unsigned"]:
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(POSITION_STATES),
                            allow_blank=True, showDropDown=False)
        dv.error = "Use one of: %s" % ", ".join(POSITION_STATES)
        dv.errorTitle = "Not an allowed decision"
        ws.add_data_validation(dv)
        dv.add("E%d:E%d" % (first, last_pos))
    pos_range = (first, last_pos)

    # Measurements
    ws = wb.create_sheet("Measurements")
    header(ws, ["Ref", "Page", "Tier", "Figure", "The sentence it supports",
                "Decision", "Source", "Note"], [12, 10, 16, 24, 80, 16, 36, 36])
    ws.append(["EXAMPLE", "0.0.0", "mrcs", "20 to 30 degrees",
               "This row shows the expected format and is ignored when the pack is read back.",
               "VERIFIED", "Author, title, edition, chapter", ""])
    for i in range(1, 9):
        c = ws.cell(row=2, column=i)
        c.font, c.fill, c.alignment = example, example_fill, wrap
    first = 3
    for f, i, g in meta["unchecked"]:
        ws.append(["%s-F%d" % (f["page_id"], i), f["page_id"], g.get("tier", ""),
                   g.get("as_written", ""), g.get("claim", ""), "", "", ""])
        for col in range(1, 9):
            ws.cell(row=ws.max_row, column=col).font = body
            ws.cell(row=ws.max_row, column=col).alignment = wrap
    last_fig = ws.max_row
    if meta["unchecked"]:
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(FIGURE_STATES),
                            allow_blank=True, showDropDown=False)
        dv.error = "Use one of: %s" % ", ".join(FIGURE_STATES)
        dv.errorTitle = "Not an allowed decision"
        ws.add_data_validation(dv)
        dv.add("F%d:F%d" % (first, last_fig))
    fig_range = (first, last_fig)

    # Progress. Formulas rather than counts, so the sheet answers "how far have I
    # got" while the reviewer is still working in the file. The ranges start below
    # the example row so it is not counted.
    ws = wb.create_sheet("Progress")
    header(ws, ["", "Recommendations", "Measurements"], [26, 18, 18])
    rows = [("Rows to decide",
             "=%d" % len(meta["unsigned"]), "=%d" % len(meta["unchecked"])),
            ("Decided",
             '=COUNTIF(Recommendations!E%d:E%d,"<>")' % pos_range,
             '=COUNTIF(Measurements!F%d:F%d,"<>")' % fig_range),
            ("Still blank", "=B2-B3", "=C2-C3")]
    for state in POSITION_STATES:
        rows.append((state,
                     '=COUNTIF(Recommendations!E%d:E%d,"%s")' % (pos_range + (state,)),
                     ""))
    for state in FIGURE_STATES:
        rows.append((state, "",
                     '=COUNTIF(Measurements!F%d:F%d,"%s")' % (fig_range + (state,))))
    for name, b, c in rows:
        ws.append([name, b, c])
        ws.cell(row=ws.max_row, column=1).font = label
        ws.cell(row=ws.max_row, column=2).font = body
        ws.cell(row=ws.max_row, column=3).font = body

    # openpyxl writes a formula with no cached value, so the Progress sheet would
    # read as blank in a previewer that does not calculate. Excel, LibreOffice and
    # Google Sheets all honour this flag and recalculate on open. LibreOffice could
    # not be used to bake the values in: it timed out in the build container, so
    # the flag is the fix rather than a convenience.
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


def read_xlsx(path):
    """Rows as the ingest side wants them: (first cell, [the cells after it])."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = []
    for name in ("Recommendations", "Measurements"):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            out.append((str(row[0]).strip(),
                        ["" if c is None else str(c).strip() for c in row[1:]]))
    return out


def write_docx(path, meta):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()

    # Landscape throughout. A recommendation is a whole sentence and a portrait
    # column makes it four lines deep, which is what makes an 85 row table
    # unreadable.
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    s.left_margin = s.right_margin = Inches(0.6)
    s.top_margin = s.bottom_margin = Inches(0.6)
    usable = s.page_width - s.left_margin - s.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for style, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
        st = doc.styles[style]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x16, 0x38, 0x2D)

    def shade(cell, hex_colour):
        from docx.oxml import OxmlElement
        el = OxmlElement("w:shd")
        el.set(qn("w:val"), "clear")
        el.set(qn("w:fill"), hex_colour)
        cell._tc.get_or_add_tcPr().append(el)

    def table(headings, widths, rows):
        """widths are fractions of the usable page width and must sum to 1."""
        t = doc.add_table(rows=1, cols=len(headings))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        for i, h in enumerate(headings):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(cell, "16382D")
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)
        # python-docx needs the width set on every cell, not on the column, or
        # Word recomputes the layout and ignores it.
        for row in t.rows:
            for i, frac in enumerate(widths):
                row.cells[i].width = int(usable * frac)
        doc.add_paragraph()
        return t

    doc.add_paragraph("Consultant review pack", style="Title")
    p = doc.add_paragraph()
    r = p.add_run("Chapter %s %s" % (meta["chapter"], meta["name"]))
    r.bold = True
    r.font.size = Pt(13)
    doc.add_paragraph(
        "Generated from the pages themselves by tools/review_pack.py. Fill the "
        "Decision column and your name; every other cell is rewritten when the "
        "pack is regenerated.")

    table(["Field", "Value"], [0.3, 0.7],
          [(k, v) for k, v in _facts_rows(meta)])

    doc.add_paragraph("What has not happened", style="Heading 1")
    doc.add_paragraph(
        "No source was retrieved for any page in this chapter. The build "
        "environment has no access to PubMed, Cochrane, a journal or a textbook, "
        "so the pipeline's fact checking stage could not run. Every page carries "
        "NOT_RUN_TOOLS_UNAVAILABLE in its verified handoff.")
    doc.add_paragraph(
        "This pack therefore stands in for the verification stage as well as for "
        "the review stage. A claim here has been checked by nobody.")

    doc.add_paragraph("The pages", style="Heading 1")
    table(["Page", "Title", "Tiers", "Words", "Published as"],
          [0.08, 0.3, 0.07, 0.08, 0.47],
          [(f["page_id"], f["title"], len(f["tiers"]), format(f["words"], ","),
            meta["published"].get(f["page_id"], "")) for f in meta["facts"]])

    doc.add_paragraph("Claims to confirm", style="Heading 1")
    doc.add_paragraph(
        "The deepest tier's key learning points on each page, which are the "
        "chapter's substantive claims. Every other statement in the body needs "
        "the same check; these are where to start.")
    for f in meta["facts"]:
        if not f["claims"]:
            continue
        doc.add_paragraph("%s %s" % (f["page_id"], f["title"]), style="Heading 2")
        for claim in f["claims"]:
            doc.add_paragraph(claim, style="List Bullet")

    doc.add_page_break()
    doc.add_paragraph("Measurements", style="Heading 1")
    if meta["unchecked"]:
        doc.add_paragraph(
            "Every figure below was written from standard teaching and none has a "
            "source against it. Allowed decisions: %s. A corrected value goes in "
            "the Note column." % ", ".join(FIGURE_STATES))
        table(["Ref", "Page", "Tier", "Figure", "The sentence it supports",
               "Decision", "Source", "Note"],
              [0.07, 0.05, 0.07, 0.12, 0.33, 0.09, 0.14, 0.13],
              [("%s-F%d" % (f["page_id"], i), f["page_id"], g.get("tier", ""),
                g.get("as_written", ""), g.get("claim", ""), "", "", "")
               for f, i, g in meta["unchecked"]])
    else:
        doc.add_paragraph("No unchecked measurement in this chapter.")

    doc.add_paragraph("Recommendations", style="Heading 1")
    if meta["unsigned"]:
        doc.add_paragraph(
            "Each sentence below tells a clinician what to do, on a page that "
            "carries no name. Allowed decisions: %s. New wording for an amended "
            "sentence goes in the Note column. The detector over collects on "
            "purpose, so a sentence that is teaching rather than clinical "
            "direction is marked NOT_CLINICAL."
            % ", ".join(POSITION_STATES))
        table(["Ref", "Page", "Tier", "Sentence", "Decision", "Reviewer", "Note"],
              [0.07, 0.05, 0.07, 0.42, 0.1, 0.13, 0.16],
              [("%s-P%d" % (f["page_id"], i), f["page_id"], p.get("tier", ""),
                p.get("as_written", ""), "", "", "")
               for f, i, p in meta["unsigned"]])
    else:
        doc.add_paragraph("No unsigned recommendation in this chapter.")

    doc.add_paragraph("How to return this", style="Heading 1")
    doc.add_paragraph(
        "Send the file back. Each decision is read into the page's own register "
        "with:")
    code = doc.add_paragraph("python3 tools/review_pack.py --ingest "
                             "docs/review/packs/chapter-%s.docx" % meta["chapter"])
    code.paragraph_format.left_indent = Inches(0.4)
    code.runs[0].font.name = "Courier New"
    doc.add_paragraph(
        "A row left blank stays outstanding and comes back in the next pack. One "
        "malformed row stops the whole file, so nothing is written half way.")
    doc.add_paragraph(
        "The style gate then enforces the result: a page carrying a "
        "recommendation that is not in its register fails POS-001, and a figure "
        "that is not in its register fails FIG-002, so a signed chapter cannot "
        "drift back to unsigned without the gate saying so.")
    doc.add_paragraph(
        "This pack does not change a page. A confirmed recommendation needs "
        "nothing; an amended one needs the prose edited, and the Note column is "
        "what the editor works from.")

    # python-docx ships a default template whose settings.xml carries
    # <w:zoom w:val="bestFit"/>. The schema requires w:percent on that element,
    # so every file python-docx writes is invalid there. Word tolerates it and
    # stricter readers do not, so it is corrected before saving.
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")

    doc.save(path)
    return path


def read_docx(path):
    from docx import Document
    doc = Document(path)
    out = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if cells and cells[0]:
                out.append((cells[0], cells[1:]))
    return out
