"""EV-2 has had reviewers and no return path since the build began. The pack
generator is that path, so these tests hold both directions of it: a pack is
generated from the work, and a filled pack is read back into the registers.

The round trip is the part worth testing. A pack that cannot be ingested is a
document, and a document is what the build already had.
"""

import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPO = os.path.dirname(ROOT)
sys.path.insert(0, os.path.join(REPO, "tools"))
import review_pack  # noqa: E402

PACKS = os.path.join(REPO, "docs", "review", "packs")


def run(*args):
    return subprocess.run([sys.executable, os.path.join(REPO, "tools", "review_pack.py")]
                          + list(args), capture_output=True, text=True, cwd=REPO)


class Generation(unittest.TestCase):

    def test_every_chapter_with_a_published_page_has_a_pack(self):
        for chapter in review_pack.chapters_with_pages():
            self.assertTrue(os.path.exists(review_pack.pack_path(chapter)),
                            "no pack for chapter %s" % chapter)

    def test_no_pack_has_drifted_from_the_work(self):
        out = run("--check")
        self.assertEqual(out.returncode, 0, out.stdout)
        self.assertNotIn("drifted", out.stdout)

    def test_a_pack_states_the_counts_its_registers_hold(self):
        chapter = "3.3"
        facts = [review_pack.page_facts(p)
                 for p in review_pack.chapter_pages(chapter)]
        unsigned = len([p for f in facts for p in f["positions"]
                        if p["verification"] == "UNSIGNED"])
        unchecked = len([g for f in facts for g in f["figures"]
                         if g["verification"] == "UNVERIFIED_FROM_MEMORY"])
        with open(review_pack.pack_path(chapter), encoding="utf-8") as fh:
            body = fh.read()
        self.assertIn("| Recommendations to sign | %d |" % unsigned, body)
        self.assertIn("| Measurements to check | %d |" % unchecked, body)

    def test_a_pack_says_verification_never_ran(self):
        """Every pack carries the statement, because every page does."""
        for chapter in review_pack.chapters_with_pages():
            with open(review_pack.pack_path(chapter), encoding="utf-8") as fh:
                self.assertIn("did not run on any page", fh.read(), chapter)

    def test_a_row_reference_survives_the_table(self):
        with open(review_pack.pack_path("3.3"), encoding="utf-8") as fh:
            rows = review_pack.parse_rows(fh.read())
        self.assertTrue(rows)
        kinds = {r[1] for r in rows}
        self.assertEqual(kinds, {"P", "F"})
        for page_id, kind, index, _ in rows:
            self.assertTrue(index >= 1)
            self.assertEqual(page_id.count("."), 2)

    def test_a_sentence_containing_a_pipe_cannot_break_a_table(self):
        self.assertNotIn("|", review_pack.cell("a | b").replace("\\|", ""))


class Ingest(unittest.TestCase):
    """The registers are real files, so each test restores what it touched."""

    def setUp(self):
        self.touched = {}
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        for path, original in self.touched.items():
            with io.open(path, "w", encoding="utf-8") as fh:
                fh.write(original)
        shutil.rmtree(self.tmp, ignore_errors=True)

    def keep(self, path):
        with open(path, encoding="utf-8") as fh:
            self.touched[path] = fh.read()

    def filled(self, chapter, replacements):
        with open(review_pack.pack_path(chapter), encoding="utf-8") as fh:
            body = fh.read()
        for old, new in replacements:
            self.assertIn(old, body)
            body = body.replace(old, new)
        path = os.path.join(self.tmp, "filled.md")
        with io.open(path, "w", encoding="utf-8") as fh:
            fh.write(body)
        return path

    def row(self, chapter, ref):
        with open(review_pack.pack_path(chapter), encoding="utf-8") as fh:
            for line in fh:
                if line.startswith("| %s |" % ref):
                    return line.rstrip("\n")
        self.fail("no row %s in chapter %s" % (ref, chapter))

    def test_a_confirmed_recommendation_reaches_its_register(self):
        reg = os.path.join(REPO, "pipeline", "config", "positions", "3.3.1.json")
        self.keep(reg)
        line = self.row("3.3", "3.3.1-P1")
        path = self.filled("3.3", [(line, line[:-4] + " CONFIRMED | Miss B Surgeon |")])
        out = run("--ingest", path)
        self.assertEqual(out.returncode, 0, out.stdout)
        with open(reg, encoding="utf-8") as fh:
            entry = json.load(fh)["positions"][0]
        self.assertEqual(entry["verification"], "CONFIRMED")
        self.assertEqual(entry["signed_off_by"], "Miss B Surgeon")

    def test_a_corrected_figure_keeps_the_replacement_wording(self):
        reg = os.path.join(REPO, "pipeline", "config", "figures", "3.3.2.json")
        self.keep(reg)
        line = self.row("3.3", "3.3.2-F1")
        path = self.filled("3.3", [(line, line[:-4] + " CORRECTED: read 0 to 25 | A textbook |")])
        out = run("--ingest", path)
        self.assertEqual(out.returncode, 0, out.stdout)
        with open(reg, encoding="utf-8") as fh:
            entry = json.load(fh)["figures"][0]
        self.assertEqual(entry["verification"], "CORRECTED")
        self.assertEqual(entry["source"], "A textbook")
        self.assertEqual(entry["note"], "read 0 to 25")

    def test_an_unfilled_pack_changes_nothing_and_is_not_an_error(self):
        """A reviewer who has read a pack and answered nothing yet has done
        nothing wrong, so ingesting it reports the blanks and succeeds."""
        reg = os.path.join(REPO, "pipeline", "config", "positions", "3.3.1.json")
        with open(reg, encoding="utf-8") as fh:
            before = fh.read()
        out = run("--ingest", review_pack.pack_path("3.3"))
        self.assertEqual(out.returncode, 0, out.stdout)
        self.assertIn("0 decision(s) written", out.stdout)
        self.assertIn("row(s) still blank", out.stdout)
        with open(reg, encoding="utf-8") as fh:
            self.assertEqual(fh.read(), before)

    def test_a_file_with_no_referenced_rows_is_refused(self):
        path = os.path.join(self.tmp, "notapack.md")
        with io.open(path, "w", encoding="utf-8") as fh:
            fh.write("# Notes\n\n| Page | Title |\n|---|---|\n| 3.3.1 | Kinematics |\n")
        out = run("--ingest", path)
        self.assertEqual(out.returncode, 1, out.stdout)
        self.assertIn("no referenced rows found", out.stdout)

    def test_an_unknown_decision_writes_nothing_at_all(self):
        """One bad row must not leave a half signed register behind."""
        reg = os.path.join(REPO, "pipeline", "config", "positions", "3.3.1.json")
        self.keep(reg)
        good = self.row("3.3", "3.3.1-P1")
        bad = self.row("3.3", "3.3.1-P2")
        path = self.filled("3.3", [(good, good[:-4] + " CONFIRMED | Miss B Surgeon |"),
                                   (bad, bad[:-4] + " LOOKS FINE | Miss B Surgeon |")])
        out = run("--ingest", path)
        self.assertEqual(out.returncode, 1, out.stdout)
        self.assertIn("nothing written", out.stdout)
        with open(reg, encoding="utf-8") as fh:
            self.assertEqual(json.load(fh)["positions"][0]["verification"], "UNSIGNED")

    def test_a_decision_survives_re_extraction(self):
        """The extractor rewrites a register from the page on every chapter, so a
        signature that does not survive it would be lost on the next build."""
        reg = os.path.join(REPO, "pipeline", "config", "positions", "3.3.1.json")
        self.keep(reg)
        line = self.row("3.3", "3.3.1-P1")
        path = self.filled("3.3", [(line, line[:-4] + " CONFIRMED | Miss B Surgeon |")])
        self.assertEqual(run("--ingest", path).returncode, 0)
        subprocess.check_call([sys.executable, os.path.join(REPO, "tools", "positions.py"),
                               "--extract"], cwd=REPO, stdout=subprocess.DEVNULL)
        with open(reg, encoding="utf-8") as fh:
            entry = json.load(fh)["positions"][0]
        self.assertEqual(entry["verification"], "CONFIRMED")
        self.assertEqual(entry["signed_off_by"], "Miss B Surgeon")


if __name__ == "__main__":
    unittest.main()


class Formats(unittest.TestCase):
    """The spreadsheet and the Word file are the formats a consultant actually
    uses, so they have to carry the same rows as the markdown and read back the
    same way. Skipped rather than failed where the two libraries are absent: a
    checkout without them can still generate, check and ingest markdown."""

    chapter = "3.3"

    @classmethod
    def setUpClass(cls):
        try:
            import openpyxl  # noqa: F401
            import docx  # noqa: F401
        except ImportError:
            raise unittest.SkipTest("openpyxl and python-docx are not installed")
        cls.tmp = tempfile.mkdtemp()
        cls.meta = review_pack.build(cls.chapter, review_pack.chapter_names())
        mod = review_pack.files_module()
        cls.xlsx = os.path.join(cls.tmp, "pack.xlsx")
        cls.docx = os.path.join(cls.tmp, "pack.docx")
        mod.write_xlsx(cls.xlsx, cls.meta)
        mod.write_docx(cls.docx, cls.meta)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def md_refs(self):
        with open(review_pack.pack_path(self.chapter), encoding="utf-8") as fh:
            return {r[0] + "-" + r[1] + str(r[2]) for r in review_pack.parse_rows(fh.read())}

    def test_all_three_formats_carry_the_same_rows(self):
        want = self.md_refs()
        self.assertTrue(want)
        for path in (self.xlsx, self.docx):
            got = {"%s-%s%d" % (r[0], r[1], r[2]) for r in review_pack.rows_from(path)}
            self.assertEqual(got, want, os.path.basename(path))

    def test_the_example_row_is_not_a_row(self):
        """Each fillable sheet carries a filled example, which must not be read
        back as a decision on a real page."""
        import openpyxl
        wb = openpyxl.load_workbook(self.xlsx)
        self.assertEqual(wb["Recommendations"]["A2"].value, "EXAMPLE")
        self.assertEqual(wb["Recommendations"]["E2"].value, "CONFIRMED")
        refs = {"%s-%s%d" % (r[0], r[1], r[2]) for r in review_pack.rows_from(self.xlsx)}
        self.assertNotIn("EXAMPLE", refs)

    def test_the_decision_columns_carry_a_dropdown(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.xlsx)
        for sheet, column, states in (
                ("Recommendations", "E", review_pack.files_module().POSITION_STATES),
                ("Measurements", "F", review_pack.files_module().FIGURE_STATES)):
            ws = wb[sheet]
            found = [dv for dv in ws.data_validations.dataValidation
                     if any(str(r).startswith(column) for r in dv.sqref.ranges)]
            self.assertTrue(found, "%s has no validation on column %s" % (sheet, column))
            for state in states:
                self.assertIn(state, found[0].formula1)

    def test_a_spreadsheet_round_trips_into_the_register(self):
        import openpyxl
        reg = os.path.join(REPO, "pipeline", "config", "positions", "3.3.1.json")
        with open(reg, encoding="utf-8") as fh:
            before = fh.read()
        filled = os.path.join(self.tmp, "filled.xlsx")
        wb = openpyxl.load_workbook(self.xlsx)
        ws = wb["Recommendations"]
        row = next(r for r in range(3, ws.max_row + 1)
                   if ws.cell(row=r, column=1).value == "3.3.1-P1")
        ws.cell(row=row, column=5).value = "AMENDED"
        ws.cell(row=row, column=6).value = "Miss C Surgeon"
        ws.cell(row=row, column=7).value = "say loading condition, not task"
        wb.save(filled)
        try:
            out = run("--ingest", filled)
            self.assertEqual(out.returncode, 0, out.stdout)
            with open(reg, encoding="utf-8") as fh:
                entry = json.load(fh)["positions"][0]
            self.assertEqual(entry["verification"], "AMENDED")
            self.assertEqual(entry["signed_off_by"], "Miss C Surgeon")
            self.assertEqual(entry["note"], "say loading condition, not task")
        finally:
            with io.open(reg, "w", encoding="utf-8") as fh:
                fh.write(before)

    def test_a_word_file_round_trips_into_the_register(self):
        import docx
        reg = os.path.join(REPO, "pipeline", "config", "figures", "3.3.2.json")
        with open(reg, encoding="utf-8") as fh:
            before = fh.read()
        filled = os.path.join(self.tmp, "filled.docx")
        d = docx.Document(self.docx)
        done = False
        for table in d.tables:
            for r in table.rows:
                if r.cells[0].text.strip() == "3.3.2-F1":
                    r.cells[5].text = "VERIFIED"
                    r.cells[6].text = "A named textbook"
                    done = True
        self.assertTrue(done, "no row 3.3.2-F1 in the Word pack")
        d.save(filled)
        try:
            out = run("--ingest", filled)
            self.assertEqual(out.returncode, 0, out.stdout)
            with open(reg, encoding="utf-8") as fh:
                entry = json.load(fh)["figures"][0]
            self.assertEqual(entry["verification"], "VERIFIED")
            self.assertEqual(entry["source"], "A named textbook")
        finally:
            with io.open(reg, "w", encoding="utf-8") as fh:
                fh.write(before)

    def test_an_unreadable_format_is_refused_by_name(self):
        path = os.path.join(self.tmp, "pack.pdf")
        with io.open(path, "w", encoding="utf-8") as fh:
            fh.write("not a pack")
        out = run("--ingest", path)
        self.assertEqual(out.returncode, 1, out.stdout)
        self.assertIn("expected .md, .xlsx or .docx", out.stdout + out.stderr)

    def test_all_generates_markdown_only(self):
        """Both binary formats embed a timestamp, so regenerating fifty of them
        on every run would churn the repository."""
        out = run("--all", "--format", "all")
        self.assertEqual(out.returncode, 2, out.stdout + out.stderr)
        self.assertIn("markdown only", out.stdout + out.stderr)

    def test_the_progress_formulas_count_the_real_rows(self):
        """LibreOffice times out in this container, so recalc.py cannot bake the
        values in. The ranges are checked directly instead: a COUNTIF over the
        wrong rows evaluates cleanly and reports the wrong number."""
        import openpyxl
        import re as _re
        wb = openpyxl.load_workbook(self.xlsx)
        for sheet, column in (("Recommendations", "E"), ("Measurements", "F")):
            ws = wb[sheet]
            rows = [r for r in range(3, ws.max_row + 1)
                    if ws.cell(row=r, column=1).value]
            self.assertTrue(rows, sheet)
            want = "%s!%s%d:%s%d" % (sheet, column, rows[0], column, rows[-1])
            found = [c.value for row in wb["Progress"].iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith("=COUNTIF")]
            self.assertTrue(any(want in f for f in found),
                            "no COUNTIF over %s; found %s" % (want, found))
        # The example row sits at row 2 and must be outside every counted range.
        for f in [c.value for row in wb["Progress"].iter_rows() for c in row
                  if isinstance(c.value, str) and c.value.startswith("=COUNTIF")]:
            first = int(_re.search(r"[A-Z](\d+):", f).group(1))
            self.assertGreaterEqual(first, 3, f)

    def test_the_workbook_recalculates_when_opened(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.xlsx)
        self.assertTrue(wb.calculation.fullCalcOnLoad)
